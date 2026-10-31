import csv
import io
import json
import os
import re
import sqlite3
import threading
import time
from concurrent.futures import ThreadPoolExecutor
from contextlib import contextmanager
import uuid
from datetime import datetime, timezone
from functools import wraps
from pathlib import Path

from flask import Flask, abort, flash, make_response, redirect, render_template, request, send_file, session, url_for
from flask_wtf.csrf import CSRFProtect
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename
from openpyxl import load_workbook, Workbook

import validator

BASE_DIR = Path(__file__).resolve().parent
INSTANCE_DIR = BASE_DIR / "instance"
INSTANCE_DIR.mkdir(exist_ok=True)
DB_PATH = Path(os.getenv("DATABASE_PATH", str(INSTANCE_DIR / "app.db")))
MAX_UPLOAD_BYTES = int(os.getenv("MAX_UPLOAD_MB", "10")) * 1024 * 1024
ALLOWED_EXTENSIONS = {"csv", "xlsx"}
RESULT_HEADERS = validator.HEADERS
EMAIL_WORKERS = max(1, int(os.getenv("EMAIL_WORKERS", "3")))
QUEUE_POLL_SECONDS = max(1, float(os.getenv("QUEUE_POLL_SECONDS", "2")))
# No fixed global job concurrency cap; one active list per user.
# Practical concurrency is limited by this PC, network, and SMTP/DNS servers.

app = Flask(__name__)
app.config.update(
    SECRET_KEY=os.getenv("SECRET_KEY", "dev-only-change-this-secret-before-deploying"),
    MAX_CONTENT_LENGTH=MAX_UPLOAD_BYTES,
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=os.getenv("COOKIE_SECURE", "false").lower() == "true",
)

csrf = CSRFProtect(app)

def now():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def connect_db():
    conn = sqlite3.connect(DB_PATH, timeout=30)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


@contextmanager
def db_session():
    db = connect_db()
    try:
        yield db
        db.commit()
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()

def init_db():
    with db_session() as db:
        db.executescript("""
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            email TEXT NOT NULL UNIQUE COLLATE NOCASE,
            password_hash TEXT NOT NULL,
            created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS jobs (
            id TEXT PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
            filename TEXT NOT NULL,
            status TEXT NOT NULL,
            total INTEGER NOT NULL DEFAULT 0,
            completed INTEGER NOT NULL DEFAULT 0,
            created_at TEXT NOT NULL,
            finished_at TEXT,
            error TEXT
        );
        CREATE TABLE IF NOT EXISTS results (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            job_id TEXT NOT NULL REFERENCES jobs(id) ON DELETE CASCADE,
            row_number INTEGER NOT NULL,
            email TEXT NOT NULL,
            data_json TEXT,
            completed_at TEXT,
            status TEXT NOT NULL DEFAULT 'pending',
            UNIQUE(job_id, row_number)
        );
        CREATE INDEX IF NOT EXISTS idx_jobs_user_created ON jobs(user_id, created_at DESC);
        CREATE INDEX IF NOT EXISTS idx_results_job_status ON results(job_id, status);
        """)
        # Migrate databases created by earlier app versions.
        result_columns = {row['name'] for row in db.execute("PRAGMA table_info(results)").fetchall()}
        if 'completed_at' not in result_columns:
            db.execute("ALTER TABLE results ADD COLUMN completed_at TEXT")
        db.execute("CREATE INDEX IF NOT EXISTS idx_results_job_completed ON results(job_id, completed_at DESC)")
        # Recover queued/running work after a restart. Results marked running
        # were interrupted mid-validation and are safe to retry.
        db.execute("UPDATE jobs SET status='queued', error=NULL WHERE status='running'")
        db.execute("UPDATE results SET status='pending' WHERE status='running'")
        db.execute("UPDATE jobs SET status='cancelled', finished_at=? WHERE status='cancelling'", (now(),))

init_db()


def login_required(fn):
    @wraps(fn)
    def wrapped(*args, **kwargs):
        if not session.get("user_id"):
            return redirect(url_for("login", next=request.path))
        # A session can outlive a deleted/invalid account; do not pass None to routes.
        if current_user() is None:
            session.clear()
            return redirect(url_for("login", next=request.path))
        return fn(*args, **kwargs)
    return wrapped


def current_user():
    if not session.get("user_id"):
        return None
    with db_session() as db:
        return db.execute("SELECT id, email FROM users WHERE id=?", (session["user_id"],)).fetchone()

@app.context_processor
def inject_user():
    return {"current_user": current_user()}


def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


def read_emails(file_storage):
    filename = secure_filename(file_storage.filename or "upload")
    ext = filename.rsplit(".", 1)[1].lower()
    emails = []
    if ext == "csv":
        raw = file_storage.stream.read()
        text = None
        for encoding in ("utf-8-sig", "utf-8", "cp1252"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise ValueError("CSV ফাইলের encoding পড়া যায়নি। UTF-8 CSV ব্যবহার করুন।")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        if not reader.fieldnames:
            raise ValueError("CSV ফাইলে header row থাকতে হবে।")
        email_col = next((h for h in reader.fieldnames if h and h.strip().lower() == "email"), None)
        if not email_col:
            raise ValueError("ফাইলে 'email' নামে একটি column থাকতে হবে।")
        for row in reader:
            value = row.get(email_col)
            if value is not None and str(value).strip():
                emails.append(str(value).strip())
    else:
        try:
            wb = load_workbook(file_storage, read_only=True, data_only=True)
        except Exception as exc:
            raise ValueError("XLSX ফাইলটি পড়া যায়নি।") from exc
        try:
            ws = wb.active
            rows = ws.iter_rows(values_only=True)
            header = next(rows, None)
            if not header:
                raise ValueError("XLSX ফাইলে header row থাকতে হবে।")
            email_index = next((i for i, h in enumerate(header) if h is not None and str(h).strip().lower() == "email"), None)
            if email_index is None:
                raise ValueError("ফাইলে 'email' নামে একটি column থাকতে হবে।")
            for row in rows:
                if email_index < len(row) and row[email_index] is not None and str(row[email_index]).strip():
                    emails.append(str(row[email_index]).strip())
        finally:
            wb.close()
    if not emails:
        raise ValueError("Email column-এ কোনো email পাওয়া যায়নি।")
    if len(emails) > 100000:
        raise ValueError("একটি ফাইলে সর্বোচ্চ 100,000টি email দেওয়া যাবে।")
    return filename, emails


def safe_export_value(value):
    # Prevent spreadsheet formula execution from untrusted imported cells.
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


def _validate_and_save(row, job_id, disposable):
    try:
        result = validator.check_email(row["email"], disposable)
    except Exception:
        result = {h: "" for h in RESULT_HEADERS}
        result["Email"] = row["email"]
        result["Status"] = "Unknown"

    with db_session() as db:
        job = db.execute("SELECT status FROM jobs WHERE id=?", (job_id,)).fetchone()
        # If cancellation was requested during this probe, discard its result.
        if not job or job["status"] in {"cancelling", "cancelled"}:
            return
        db.execute(
            "UPDATE results SET status='done', data_json=?, completed_at=? WHERE id=? AND job_id=?",
            (json.dumps(result, ensure_ascii=False), datetime.now(timezone.utc).isoformat(timespec="microseconds"), row["id"], job_id),
        )
        db.execute("UPDATE jobs SET completed=completed+1 WHERE id=?", (job_id,))


def run_job(job_id):
    """Process a queued job with bounded parallelism; results remain in SQLite."""
    try:
        with db_session() as db:
            job = db.execute("SELECT status FROM jobs WHERE id=?", (job_id,)).fetchone()
            if not job or job["status"] not in {"queued", "running"}:
                return
            db.execute("UPDATE jobs SET status='running', error=NULL WHERE id=?", (job_id,))

        disposable = validator.load_disposable_domains()
        with ThreadPoolExecutor(max_workers=EMAIL_WORKERS, thread_name_prefix="email-check") as pool:
            while True:
                with db_session() as db:
                    job = db.execute("SELECT status FROM jobs WHERE id=?", (job_id,)).fetchone()
                    if not job:
                        return
                    if job["status"] in {"cancelling", "cancelled"}:
                        db.execute("UPDATE results SET status='pending' WHERE job_id=? AND status='running'", (job_id,))
                        db.execute("UPDATE jobs SET status='cancelled', finished_at=? WHERE id=?", (now(), job_id))
                        break
                    rows = db.execute(
                        "SELECT id, email FROM results WHERE job_id=? AND status='pending' ORDER BY row_number LIMIT 100",
                        (job_id,),
                    ).fetchall()
                    if not rows:
                        db.execute(
                            "UPDATE jobs SET status='completed', finished_at=?, error=NULL WHERE id=? AND status='running'",
                            (now(), job_id),
                        )
                        break
                    db.executemany(
                        "UPDATE results SET status='running' WHERE id=? AND status='pending'",
                        [(row["id"],) for row in rows],
                    )

                futures = [pool.submit(_validate_and_save, row, job_id, disposable) for row in rows]
                # Stop queued probes when cancellation is requested; in-flight network calls finish first.
                for index, future in enumerate(futures):
                    with db_session() as db:
                        status_row = db.execute("SELECT status FROM jobs WHERE id=?", (job_id,)).fetchone()
                    if not status_row or status_row["status"] in {"cancelling", "cancelled"}:
                        for pending_future in futures[index:]:
                            pending_future.cancel()
                        break
                    future.result()
                # Loop back to finalize cancellation or fetch the next batch.
    except Exception as exc:
        with db_session() as db:
            db.execute(
                "UPDATE results SET status='pending' WHERE job_id=? AND status='running'",
                (job_id,),
            )
            db.execute("UPDATE jobs SET status='failed', error=? WHERE id=?", (str(exc)[:500], job_id))


def queue_worker():
    """Run one list per user concurrently, without a fixed global job cap."""
    active_user_ids = set()
    active_lock = threading.Lock()

    def run_user_job(job_id, user_id):
        try:
            run_job(job_id)
        finally:
            with active_lock:
                active_user_ids.discard(user_id)

    while True:
        launched = False
        try:
            with active_lock:
                active_ids = set(active_user_ids)
            with db_session() as db:
                rows = db.execute(
                    "SELECT id, user_id FROM jobs WHERE status='queued' ORDER BY created_at"
                ).fetchall()
            # Launch the oldest queued job for each user not already running.
            # Other jobs for the same user remain queued until this one finishes.
            selected = []
            selected_users = set(active_ids)
            for row in rows:
                uid = row["user_id"]
                if uid in selected_users:
                    continue
                selected.append((row["id"], uid))
                selected_users.add(uid)
            for job_id, user_id in selected:
                with active_lock:
                    if user_id in active_user_ids:
                        continue
                    active_user_ids.add(user_id)
                thread = threading.Thread(
                    target=run_user_job,
                    args=(job_id, user_id),
                    name=f"list-job-{job_id}",
                    daemon=True,
                )
                thread.start()
                launched = True
            if not launched:
                time.sleep(QUEUE_POLL_SECONDS)
        except Exception as exc:
            app.logger.exception("Background queue error: %s", exc)
            time.sleep(QUEUE_POLL_SECONDS)


_queue_thread = threading.Thread(target=queue_worker, name="sqlite-job-queue", daemon=True)
_queue_thread.start()


def user_job(job_id):
    user = current_user()
    if not user:
        abort(401)
    with db_session() as db:
        job = db.execute("SELECT * FROM jobs WHERE id=? AND user_id=?", (job_id, user["id"])).fetchone()
    if not job:
        abort(404)
    return job

@app.route("/")
def index():
    if session.get("user_id"):
        return redirect(url_for("dashboard"))
    return render_template("landing.html")

@app.route("/signup", methods=["GET", "POST"])
def signup():
    if request.method == "POST":
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")
        if not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", email):
            flash("একটি সঠিক email address দিন।", "error")
        elif len(password) < 10:
            flash("Password কমপক্ষে ১০ অক্ষরের হতে হবে।", "error")
        else:
            try:
                with db_session() as db:
                    cur = db.execute("INSERT INTO users(email,password_hash,created_at) VALUES(?,?,?)", (email, generate_password_hash(password), now()))
                    session.clear()
                    session["user_id"] = cur.lastrowid
                return redirect(url_for("dashboard"))
            except sqlite3.IntegrityError:
                flash("এই email দিয়ে account আগে থেকেই আছে।", "error")
    return render_template("auth.html", mode="signup")

@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")
        with db_session() as db:
            user = db.execute("SELECT * FROM users WHERE email=?", (email,)).fetchone()
        if user and check_password_hash(user["password_hash"], password):
            session.clear()
            session["user_id"] = user["id"]
            return redirect(url_for("dashboard"))
        flash("Email অথবা password সঠিক নয়।", "error")
    return render_template("auth.html", mode="login")

@app.route("/logout", methods=["POST"])
@login_required
def logout():
    session.clear()
    return redirect(url_for("index"))

@app.route("/dashboard")
@login_required
def dashboard():
    user = current_user()
    if user is None:
        session.clear()
        return redirect(url_for("login", next=request.path))
    with db_session() as db:
        jobs = db.execute("SELECT * FROM jobs WHERE user_id=? ORDER BY created_at DESC LIMIT 50", (user["id"],)).fetchall()
        totals = db.execute("SELECT COUNT(*) AS jobs, COALESCE(SUM(total),0) AS emails FROM jobs WHERE user_id=?", (user["id"],)).fetchone()
    return render_template("dashboard.html", jobs=jobs, totals=totals)

@app.route("/validate", methods=["GET", "POST"])
@login_required
def validate_upload():
    if request.method == "POST":
        uploaded = request.files.get("file")
        if not uploaded or not uploaded.filename:
            flash("CSV বা XLSX ফাইল নির্বাচন করুন।", "error")
            return redirect(url_for("validate_upload"))
        if not allowed_file(uploaded.filename):
            flash("শুধু CSV এবং XLSX ফাইল গ্রহণ করা হয়।", "error")
            return redirect(url_for("validate_upload"))
        try:
            filename, emails = read_emails(uploaded)
        except ValueError as exc:
            flash(str(exc), "error")
            return redirect(url_for("validate_upload"))
        user = current_user()
        job_id = uuid.uuid4().hex
        with db_session() as db:
            # Queue immediately after upload. The worker starts jobs automatically;
            # per-user sequencing is enforced by queue_worker().
            db.execute("INSERT INTO jobs(id,user_id,filename,status,total,created_at) VALUES(?,?,?,?,?,?)", (job_id, user["id"], filename, "queued", len(emails), now()))
            db.executemany("INSERT INTO results(job_id,row_number,email) VALUES(?,?,?)", [(job_id, i, email) for i, email in enumerate(emails, 1)])
        return redirect(url_for("job_detail", job_id=job_id))
    return render_template("validate.html")

def job_summary(job_id):
    """Summarize finished rows only; percentages are of validated results."""
    counts = {"Likely Valid": 0, "Catch All": 0, "Unknown": 0, "Invalid": 0}
    with db_session() as db:
        rows = db.execute(
            "SELECT data_json FROM results WHERE job_id=? AND status='done'",
            (job_id,),
        ).fetchall()
    for row in rows:
        try:
            data = json.loads(row["data_json"] or "{}")
            smtp_result = str(data.get("SMTP result") or "")
            status = (data.get("Status") or "Unknown").strip()
            if smtp_result.lower().startswith("likely catch-all") or status.lower() == "catch all":
                status = "Catch All"
            elif status.lower() in {"valid", "likely valid"}:
                # Legacy generic Valid is not evidence of catch-all.
                status = "Likely Valid"
        except (ValueError, TypeError):
            status = "Unknown"
        if status not in counts:
            status = "Unknown"
        counts[status] += 1
    completed = sum(counts.values())
    percentages = {key: round((value / completed) * 100, 1) if completed else 0 for key, value in counts.items()}
    return {"counts": counts, "percentages": percentages, "completed_results": completed}


@app.route("/jobs/<job_id>/delete", methods=["POST"])
@login_required
def delete_job(job_id):
    """Delete one saved validation list and its results for the signed-in user."""
    job = user_job(job_id)
    if job["status"] in {"running", "cancelling"}:
        flash("চলমান validation আগে Cancel করুন। শেষ হলে list delete করতে পারবেন।", "error")
        return redirect(url_for("dashboard"))
    with db_session() as db:
        # Ownership is checked again in the DELETE itself. Results are removed by
        # the foreign-key cascade, and other users' jobs are never affected.
        cur = db.execute(
            "DELETE FROM jobs WHERE id=? AND user_id=? AND status NOT IN ('running','cancelling')",
            (job_id, session["user_id"]),
        )
        if cur.rowcount:
            flash(f'"{job["filename"]}" list এবং তার saved results delete হয়েছে।', "success")
        else:
            flash("List delete করা যায়নি। Validation চলমান থাকতে পারে।", "error")
    return redirect(url_for("dashboard"))


@app.route("/jobs/<job_id>")
@login_required
def job_detail(job_id):
    job = user_job(job_id)
    with db_session() as db:
        total_results = db.execute(
            "SELECT COUNT(*) FROM results WHERE job_id=? AND status='done'", (job_id,)
        ).fetchone()[0]
        page_size = 10
        total_pages = max(1, (total_results + page_size - 1) // page_size)
        initial_page = 1
        offset = 0
        rows = db.execute(
            "SELECT row_number,email,data_json,status FROM results WHERE job_id=? AND status='done' ORDER BY completed_at ASC, id ASC LIMIT ? OFFSET ?",
            (job_id, page_size, offset),
        ).fetchall()
    preview = []
    for row in rows:
        data = json.loads(row["data_json"]) if row["data_json"] else {"Email": row["email"], "Status": row["status"].title()}
        preview.append(data)
    summary = job_summary(job_id)
    return render_template("job.html", job=job, preview=preview, headers=RESULT_HEADERS, summary=summary,
                           initial_page=initial_page, total_pages=total_pages, total_results=total_results)

@app.route("/jobs/<job_id>/start", methods=["POST"])
@login_required
def start_job(job_id):
    job = user_job(job_id)
    if job["status"] == "ready":
        with db_session() as db:
            db.execute(
                "UPDATE jobs SET status='queued' WHERE id=? AND user_id=? AND status='ready'",
                (job_id, session["user_id"]),
            )
    return redirect(url_for("job_detail", job_id=job_id))

@app.route("/jobs/<job_id>/cancel", methods=["POST"])
@login_required
def cancel_job(job_id):
    job = user_job(job_id)
    with db_session() as db:
        if job["status"] == "queued":
            # A queued job has not started probing yet, so cancellation is immediate.
            db.execute("UPDATE jobs SET status='cancelled', finished_at=? WHERE id=? AND user_id=? AND status='queued'", (now(), job_id, session["user_id"]))
        elif job["status"] == "running":
            # Active SMTP/DNS calls cannot be forcibly interrupted; workers stop after in-flight calls return.
            db.execute("UPDATE jobs SET status='cancelling' WHERE id=? AND user_id=? AND status='running'", (job_id, session["user_id"]))
    return redirect(url_for("job_detail", job_id=job_id))

@app.route("/api/jobs/<job_id>")
@login_required
def job_status(job_id):
    job = user_job(job_id)
    progress = int((job["completed"] / job["total"]) * 100) if job["total"] else 0
    summary = job_summary(job["id"])
    page_size = 10
    requested_page = request.args.get("page", type=int) or 1
    with db_session() as db:
        total_results = db.execute(
            "SELECT COUNT(*) FROM results WHERE job_id=? AND status='done'", (job_id,)
        ).fetchone()[0]
        total_pages = max(1, (total_results + page_size - 1) // page_size)
        page = max(1, min(requested_page, total_pages))
        offset = (page - 1) * page_size
        result_rows = db.execute(
            "SELECT email, data_json FROM results WHERE job_id=? AND status='done' ORDER BY completed_at ASC, id ASC LIMIT ? OFFSET ?",
            (job_id, page_size, offset),
        ).fetchall()
    page_results = []
    for row in result_rows:
        try:
            result = json.loads(row["data_json"] or "{}")
        except (ValueError, TypeError):
            result = {}
        if not result:
            result = {"Email": row["email"], "Status": "Unknown"}
        smtp_result = str(result.get("SMTP result") or "")
        if smtp_result.lower().startswith("likely catch-all") or str(result.get("Status") or "").lower() == "catch all":
            result["Status"] = "Catch All"
        elif str(result.get("Status") or "").lower() in {"valid", "likely valid"}:
            result["Status"] = "Likely Valid"
        page_results.append(result)
    return {"id": job["id"], "status": job["status"], "total": job["total"], "completed": job["completed"], "progress": progress, "error": job["error"], "headers": RESULT_HEADERS, "page_results": page_results, "page": page, "page_size": page_size, "total_pages": total_pages, "total_results": total_results, **summary}

@app.route("/jobs/<job_id>/export/<fmt>")
@login_required
def export_job(job_id, fmt):
    job = user_job(job_id)
    if job["status"] not in {"completed", "cancelled", "failed", "interrupted"}:
        flash("Validation চলাকালীন export করতে চাইলে আগে Cancel করুন।", "error")
        return redirect(url_for("job_detail", job_id=job_id))
    if fmt not in {"csv", "xlsx"}:
        abort(404)
    with db_session() as db:
        rows = db.execute("SELECT data_json,email FROM results WHERE job_id=? AND status='done' ORDER BY row_number", (job_id,)).fetchall()
    records = []
    for row in rows:
        data = json.loads(row["data_json"]) if row["data_json"] else {"Email": row["email"], "Status": "Unknown"}
        smtp_result = str(data.get("SMTP result") or "")
        if smtp_result.lower().startswith("likely catch-all") or str(data.get("Status") or "").lower() == "catch all":
            data["Status"] = "Catch All"
        elif str(data.get("Status") or "").lower() in {"valid", "likely valid"}:
            data["Status"] = "Likely Valid"
        records.append([safe_export_value(data.get(header, "")) for header in RESULT_HEADERS])
    if not records:
        flash("Export করার মতো completed result এখনো নেই।", "error")
        return redirect(url_for("job_detail", job_id=job_id))
    base = Path(job["filename"]).stem[:80] or "emails"
    if fmt == "csv":
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(RESULT_HEADERS)
        writer.writerows(records)
        response = make_response(output.getvalue().encode("utf-8-sig"))
        response.headers["Content-Type"] = "text/csv; charset=utf-8"
        response.headers["Content-Disposition"] = f'attachment; filename="{base}_validated.csv"'
        return response
    wb = Workbook()
    ws = wb.active
    ws.title = "Email Results"
    ws.append(RESULT_HEADERS)
    for record in records:
        ws.append(record)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        letter = col[0].column_letter
        ws.column_dimensions[letter].width = min(max(max(len(str(c.value or "")) for c in col) + 2, 12), 42)
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return send_file(stream, as_attachment=True, download_name=f"{base}_validated.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.errorhandler(413)
def too_large(_):
    flash(f"ফাইলটি অনেক বড়। সর্বোচ্চ {MAX_UPLOAD_BYTES // (1024*1024)} MB upload করা যাবে।", "error")
    return redirect(url_for("validate_upload"))


if __name__ == "__main__":
    from waitress import serve

    serve(
        app,
        host="127.0.0.1",
        port=int(os.getenv("PORT", "5000"))
    )

